"""Erzeugt finanzen/Finanzmodell.xlsx mit Formeln aus den Annahmen in modell.py."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from modell import GEMEINSAM, SZENARIEN, MONATE

F = "Arial"
BLUE, GREEN = Font(name=F, color="0000FF"), Font(name=F, color="008000")
NORM, BOLD = Font(name=F), Font(name=F, bold=True)
H1 = Font(name=F, bold=True, size=14)
HEAD = PatternFill("solid", fgColor="0F5C4A"); HEADF = Font(name=F, bold=True, color="FFFFFF")
KEY = PatternFill("solid", fgColor="FFFF00")
SUB = PatternFill("solid", fgColor="EFE8DB")
EUR = '#,##0 €;(#,##0 €);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
TOP = Border(top=Side(style="thin"))

wb = Workbook()
A = wb.active; A.title = "Annahmen"
A["A1"] = "Finanzmodell Agentur – Annahmen"; A["A1"].font = H1
A["A2"] = "Blaue Zahlen sind Eingaben und dürfen geändert werden. Gelb = wichtigste Hebel. Alle anderen Blätter rechnen automatisch."; A["A2"].font = Font(name=F, italic=True)

labels = [
 ("preis_web", "Ø Einmalpreis Website", EUR, "Mix aus Website Start (1.490 €) und Wachstum (2.990 €), Preisliste der Website"),
 ("preis_pflege", "Ø Pflege & Hosting pro Monat", EUR, "Mix aus 49 € und 89 €"),
 ("preis_seo", "Ø SEO-Retainer pro Monat", EUR, "Mix aus SEO Lokal (390 €) und SEO Plus (690 €)"),
 ("preis_ads", "Ads-Betreuung pro Monat", EUR, "Preisliste"),
 ("preis_ads_setup", "Ads-Einrichtung einmalig", EUR, "Preisliste"),
 ("preis_programm", "Wachstumsprogramm pro Monat", EUR, "Preisliste, 12 Monate Laufzeit"),
 ("churn_pflege", "Kündigungsquote Pflege / Monat", PCT, "Annahme: Websites werden selten gewechselt"),
 ("churn_seo", "Kündigungsquote SEO / Monat", PCT, "Annahme, nach 6 Monaten Mindestlaufzeit"),
 ("churn_ads", "Kündigungsquote Ads / Monat", PCT, "Annahme: monatlich kündbar"),
 ("churn_programm", "Kündigungsquote Programm / Monat", PCT, "Annahme: 12-Monats-Verträge, Verlängerungsquote ca. 75 %"),
 ("fix_kosten", "Fixkosten pro Monat", EUR, "Tools 100 €, Infrastruktur 30 €, Haftpflicht 25 €, Buchhaltung 30 €, Domain/Mail 20 €"),
 ("marketing", "Eigenes Marketing pro Monat", EUR, "Ads/Tools für die eigene Kundengewinnung"),
 ("std_satz", "Stundensatz Freelancer", EUR, "Zugekaufte Stunden (Text, Design, Entwicklung)"),
 ("zukauf_web_std", "Zukauf-Stunden je neuer Website", NUM1, ""),
 ("zukauf_seo_std", "Zukauf-Stunden je SEO-Kunde / Monat", NUM1, ""),
 ("zukauf_programm_std", "Zukauf-Stunden je Programm-Kunde / Monat", NUM1, ""),
 ("eigen_web_std", "Eigene Stunden je neuer Website", NUM1, "Inkl. Vertrieb, Abstimmung, Qualitätskontrolle"),
 ("eigen_seo_std", "Eigene Stunden je SEO-Kunde / Monat", NUM1, ""),
 ("eigen_ads_std", "Eigene Stunden je Ads-Kunde / Monat", NUM1, ""),
 ("eigen_programm_std", "Eigene Stunden je Programm-Kunde / Monat", NUM1, ""),
 ("eigen_pflege_std", "Eigene Stunden je Pflege-Kunde / Monat", NUM1, "Updates laufen weitgehend automatisiert"),
]
REF = {}
A["A4"], A["B4"], A["C4"] = "Gemeinsame Annahmen", "Wert", "Quelle / Begründung"
for c in "ABC":
    A[f"{c}4"].font, A[f"{c}4"].fill = HEADF, HEAD
for i, (k, lab, fmt, src) in enumerate(labels, start=5):
    A[f"A{i}"] = lab; A[f"A{i}"].font = NORM
    A[f"B{i}"] = GEMEINSAM[k]; A[f"B{i}"].font = BLUE; A[f"B{i}"].number_format = fmt
    A[f"C{i}"] = src; A[f"C{i}"].font = NORM
    REF[k] = f"Annahmen!$B${i}"
for k in ("preis_seo", "churn_seo", "std_satz"):
    A[REF[k].split("!")[1].replace("$", "")].fill = KEY

r0 = 5 + len(labels) + 2
A[f"A{r0}"] = "Szenario-Hebel"
names = list(SZENARIEN)
for j, n in enumerate(names):
    A.cell(r0, 2 + j, n)
A.cell(r0, 5, "Erläuterung")
for c in range(1, 6):
    A.cell(r0, c).font, A.cell(r0, c).fill = HEADF, HEAD
hebel = [("web", 0, "Neue Website-Kunden / Monat, Monat 1–6", NUM1), ("web", 1, "Neue Website-Kunden / Monat, Monat 7–12", NUM1), ("web", 2, "Neue Website-Kunden / Monat, Monat 13–24", NUM1),
         ("q_seo", None, "Anteil Website-Kunden mit SEO-Retainer", PCT), ("q_ads", None, "Anteil Website-Kunden mit Ads-Betreuung", PCT),
         ("prog", 0, "Neue Programm-Kunden / Monat, Monat 1–6", NUM1), ("prog", 1, "Neue Programm-Kunden / Monat, Monat 7–12", NUM1), ("prog", 2, "Neue Programm-Kunden / Monat, Monat 13–24", NUM1)]
HREF = {n: {} for n in names}
for i, (k, idx, lab, fmt) in enumerate(hebel, start=r0 + 1):
    A[f"A{i}"] = lab; A[f"A{i}"].font = NORM
    for j, n in enumerate(names):
        v = SZENARIEN[n][k] if idx is None else SZENARIEN[n][k][idx]
        c = A.cell(i, 2 + j, v); c.font = BLUE; c.number_format = fmt; c.fill = KEY
        HREF[n][(k, idx)] = f"Annahmen!${L(2 + j)}${i}"
A.cell(r0 + 1, 5, "Dezimalwerte = Durchschnitt, z. B. 1,5 = drei Kunden in zwei Monaten").font = NORM
A.cell(r0 + 6, 5, "Programm-Kunden kommen zusätzlich zu den Website-Kunden").font = NORM
A.column_dimensions["A"].width = 46; A.column_dimensions["B"].width = 14; A.column_dimensions["C"].width = 14
A.column_dimensions["D"].width = 14; A.column_dimensions["E"].width = 60
A.freeze_panes = "A5"

ROWS = {}
def build(n):
    ws = wb.create_sheet(n)
    h = HREF[n]
    ws["A1"] = f"Szenario {n} – Monatsplanung (24 Monate)"; ws["A1"].font = H1
    ws["A2"] = "Alle Werte werden aus dem Blatt „Annahmen“ berechnet. Kundenzahlen sind Erwartungswerte (Durchschnitte)."; ws["A2"].font = Font(name=F, italic=True)
    ws["A3"] = "Monat"
    for m in range(1, MONATE + 1):
        c = ws.cell(3, m + 1, m)
    for c in range(1, MONATE + 2):
        ws.cell(3, c).font, ws.cell(3, c).fill = HEADF, HEAD
    lines = [
     ("Neue Kunden", None, None),
     ("Neue Websites", "web", lambda c, p: f'=IF({c}$3<=6,{h[("web",0)]},IF({c}$3<=12,{h[("web",1)]},{h[("web",2)]}))'),
     ("Neue SEO-Retainer", "n_seo", lambda c, p: f"={c}{R['web']}*{h[('q_seo',None)]}"),
     ("Neue Ads-Betreuungen", "n_ads", lambda c, p: f"={c}{R['web']}*{h[('q_ads',None)]}"),
     ("Neue Programm-Kunden", "n_prog", lambda c, p: f'=IF({c}$3<=6,{h[("prog",0)]},IF({c}$3<=12,{h[("prog",1)]},{h[("prog",2)]}))'),
     ("Aktive Kunden (Monatsende)", None, None),
     ("Pflege & Hosting", "k_pflege", lambda c, p: f"={'0' if p is None else p+str(R['k_pflege'])}*(1-{REF['churn_pflege']})+{c}{R['web']}"),
     ("SEO-Retainer", "k_seo", lambda c, p: f"={'0' if p is None else p+str(R['k_seo'])}*(1-{REF['churn_seo']})+{c}{R['n_seo']}"),
     ("Ads-Betreuung", "k_ads", lambda c, p: f"={'0' if p is None else p+str(R['k_ads'])}*(1-{REF['churn_ads']})+{c}{R['n_ads']}"),
     ("Wachstumsprogramm", "k_prog", lambda c, p: f"={'0' if p is None else p+str(R['k_prog'])}*(1-{REF['churn_programm']})+{c}{R['n_prog']}"),
     ("Umsatz (€)", None, None),
     ("MRR Pflege", "m_pflege", lambda c, p: f"={c}{R['k_pflege']}*{REF['preis_pflege']}"),
     ("MRR SEO", "m_seo", lambda c, p: f"={c}{R['k_seo']}*{REF['preis_seo']}"),
     ("MRR Ads", "m_ads", lambda c, p: f"={c}{R['k_ads']}*{REF['preis_ads']}"),
     ("MRR Programm", "m_prog", lambda c, p: f"={c}{R['k_prog']}*{REF['preis_programm']}"),
     ("MRR gesamt", "mrr", lambda c, p: f"=SUM({c}{R['m_pflege']}:{c}{R['m_prog']})"),
     ("Einmalumsatz (Websites, Ads-Einrichtung)", "einmal", lambda c, p: f"={c}{R['web']}*{REF['preis_web']}+{c}{R['n_ads']}*{REF['preis_ads_setup']}"),
     ("Umsatz gesamt", "umsatz", lambda c, p: f"={c}{R['mrr']}+{c}{R['einmal']}"),
     ("Kosten (€)", None, None),
     ("Fixkosten", "fix", lambda c, p: f"={REF['fix_kosten']}"),
     ("Eigenes Marketing", "mkt", lambda c, p: f"={REF['marketing']}"),
     ("Zugekaufte Stunden", "z_std", lambda c, p: f"={c}{R['web']}*{REF['zukauf_web_std']}+{c}{R['k_seo']}*{REF['zukauf_seo_std']}+{c}{R['k_prog']}*{REF['zukauf_programm_std']}"),
     ("Kosten zugekaufte Stunden", "z_eur", lambda c, p: f"={c}{R['z_std']}*{REF['std_satz']}"),
     ("Kosten gesamt", "kosten", lambda c, p: f"={c}{R['fix']}+{c}{R['mkt']}+{c}{R['z_eur']}"),
     ("Ergebnis", None, None),
     ("Ergebnis vor Steuern und Unternehmerlohn", "erg", lambda c, p: f"={c}{R['umsatz']}-{c}{R['kosten']}"),
     ("Kasse kumuliert", "kasse", lambda c, p: f"={'0' if p is None else p+str(R['kasse'])}+{c}{R['erg']}"),
     ("Eigene Arbeitsstunden im Monat", "eigen", lambda c, p: f"={c}{R['web']}*{REF['eigen_web_std']}+{c}{R['k_seo']}*{REF['eigen_seo_std']}+{c}{R['k_ads']}*{REF['eigen_ads_std']}+{c}{R['k_prog']}*{REF['eigen_programm_std']}+{c}{R['k_pflege']}*{REF['eigen_pflege_std']}"),
    ]
    R = {}
    row = 4
    for lab, key, _ in lines:
        row += 1
        if key: R[key] = row
        else: row += 0
    row = 4
    for lab, key, f in lines:
        row += 1
        ws.cell(row, 1, lab).font = BOLD if key is None or key in ("mrr", "umsatz", "kosten", "erg", "kasse") else NORM
        if key is None:
            for c in range(1, MONATE + 2): ws.cell(row, c).fill = SUB
            continue
        fmt = NUM1 if key in ("web", "n_seo", "n_ads", "n_prog", "k_pflege", "k_seo", "k_ads", "k_prog", "z_std", "eigen") else EUR
        for m in range(1, MONATE + 1):
            col = L(m + 1); prev = L(m) if m > 1 else None
            c = ws.cell(row, m + 1, f(col, prev)); c.number_format = fmt
            c.font = GREEN if key in ("web", "n_prog", "fix", "mkt") else (BOLD if key in ("mrr", "umsatz", "kosten", "erg", "kasse") else NORM)
            if key in ("mrr", "umsatz", "kosten", "erg"): c.border = TOP
    ws.column_dimensions["A"].width = 42
    for m in range(1, MONATE + 1): ws.column_dimensions[L(m + 1)].width = 10
    ws.freeze_panes = "B4"
    ROWS[n] = R

for n in names: build(n)

U = wb.create_sheet("Übersicht", 0)
U["A1"] = "Übersicht – drei Szenarien"; U["A1"].font = H1
U["A2"] = "Ergebnis vor Steuern und vor Unternehmerlohn. Startkapital: keins, alle Kosten werden aus dem laufenden Umsatz gedeckt."; U["A2"].font = Font(name=F, italic=True)
U["A4"] = "Kennzahl"
for j, n in enumerate(names): U.cell(4, 2 + j, n)
for c in range(1, 5): U.cell(4, c).font, U.cell(4, c).fill = HEADF, HEAD
last = L(MONATE + 1)
kpis = [("MRR Monat 6", lambda n, R: f"='{n}'!G{R['mrr']}", EUR), ("MRR Monat 12", lambda n, R: f"='{n}'!M{R['mrr']}", EUR),
        ("MRR Monat 18", lambda n, R: f"='{n}'!S{R['mrr']}", EUR), ("MRR Monat 24", lambda n, R: f"='{n}'!{last}{R['mrr']}", EUR),
        ("Monat, in dem 10.000 € MRR erreicht werden", lambda n, R: f"=IF(MAX('{n}'!B{R['mrr']}:{last}{R['mrr']})>=10000,COUNTIF('{n}'!B{R['mrr']}:{last}{R['mrr']},\"<10000\")+1,\"nicht in 24 Monaten\")", "0"),
        ("Umsatz Jahr 1", lambda n, R: f"=SUM('{n}'!B{R['umsatz']}:M{R['umsatz']})", EUR), ("Umsatz Jahr 2", lambda n, R: f"=SUM('{n}'!N{R['umsatz']}:{last}{R['umsatz']})", EUR),
        ("Ergebnis Jahr 1", lambda n, R: f"=SUM('{n}'!B{R['erg']}:M{R['erg']})", EUR), ("Ergebnis Jahr 2", lambda n, R: f"=SUM('{n}'!N{R['erg']}:{last}{R['erg']})", EUR),
        ("Kapitalbedarf (tiefster Kassenstand)", lambda n, R: f"=MIN(0,MIN('{n}'!B{R['kasse']}:{last}{R['kasse']}))", EUR),
        ("Aktive Kunden Monat 24 (Pflege)", lambda n, R: f"='{n}'!{last}{R['k_pflege']}", NUM1),
        ("Eigene Arbeitsstunden Monat 24", lambda n, R: f"='{n}'!{last}{R['eigen']}", NUM1)]
for i, (lab, f, fmt) in enumerate(kpis, start=5):
    U.cell(i, 1, lab).font = NORM
    for j, n in enumerate(names):
        c = U.cell(i, 2 + j, f(n, ROWS[n])); c.font = GREEN; c.number_format = fmt
U.column_dimensions["A"].width = 46
for c in "BCD": U.column_dimensions[c].width = 18

# Diagrammdaten (MRR je Monat) + Linien
d0 = 5 + len(kpis) + 2
U.cell(d0, 1, "MRR je Monat").font = BOLD
for j, n in enumerate(names): U.cell(d0, 2 + j, n).font = BOLD
for m in range(1, MONATE + 1):
    U.cell(d0 + m, 1, m).font = NORM
    for j, n in enumerate(names):
        c = U.cell(d0 + m, 2 + j, f"='{n}'!{L(m + 1)}{ROWS[n]['mrr']}"); c.font = GREEN; c.number_format = EUR
ch = LineChart(); ch.title = "Monatlich wiederkehrender Umsatz (MRR)"; ch.y_axis.title = "€"; ch.x_axis.title = "Monat"
ch.add_data(Reference(U, min_col=2, max_col=4, min_row=d0, max_row=d0 + MONATE), titles_from_data=True)
ch.set_categories(Reference(U, min_col=1, min_row=d0 + 1, max_row=d0 + MONATE))
for s, col in zip(ch.series, ("B0602A", "0F5C4A", "6B8F84")):
    s.graphicalProperties.line.solidFill = col; s.graphicalProperties.line.width = 28000; s.smooth = False
ch.height, ch.width = 9, 18
U.add_chart(ch, "F4")

wb.save("Finanzmodell.xlsx")
print("gespeichert")
